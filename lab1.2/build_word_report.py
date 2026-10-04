"""Сборка итогового Word-отчёта по лабораторной работе 1.2.

Скрипт не выполняет статистические расчёты. Он читает готовые CSV и PNG из
каталога results и оформляет их по образцу отчёта лабораторной работы 1.1.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
RESULTS = BASE / "results"
REFERENCE = BASE / ".tmp" / "word_report" / "reference.docx"
OUTPUT = BASE / "Отчёт_ЛР1.2_Регрессионный_анализ.docx"

FONT = "Times New Roman"
TEXT_WIDTH_DXA = 9354


def set_run_font(run, size=14, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    fmt = normal.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.first_line_indent = Cm(1.25)
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(10 if name == "Heading 1" else 6)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.0
        pf.keep_with_next = True

    if "Equation" not in [style.name for style in doc.styles]:
        equation = doc.styles.add_style("Equation", 1)
    else:
        equation = doc.styles["Equation"]
    equation.font.name = FONT
    equation.font.size = Pt(12)
    equation._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    equation._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    epf = equation.paragraph_format
    epf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    epf.first_line_indent = Cm(0)
    epf.space_before = Pt(4)
    epf.space_after = Pt(4)
    epf.line_spacing = 1.0
    epf.keep_together = True


def replace_title_text(doc):
    for paragraph in doc.paragraphs:
        if "ЛАБОРАТОРНАЯ РАБОТА №1" in paragraph.text:
            paragraph.clear()
            run = paragraph.add_run("ЛАБОРАТОРНАЯ РАБОТА №1.2")
            set_run_font(run, 14, bold=True)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif paragraph.text.strip() == "по дисциплине «Методы анализа данных»":
            paragraph.clear()
            first = paragraph.add_run("«Регрессионный анализ»")
            set_run_font(first, 14, bold=True)
            first.add_break()
            second = paragraph.add_run("по дисциплине «Методы анализа данных»")
            set_run_font(second, 12)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def first_section_break_child(doc):
    body = doc._element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p") and child.find(".//w:sectPr", namespaces=child.nsmap) is not None:
            return child
    raise RuntimeError("В образце не найден разрыв первого раздела")


def clear_after_title_page(doc):
    body = doc._element.body
    break_child = first_section_break_child(doc)
    passed_break = False
    for child in list(body.iterchildren()):
        if child is break_child:
            passed_break = True
            continue
        if passed_break and child.tag != qn("w:sectPr"):
            body.remove(child)


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_toc(doc):
    heading = doc.add_paragraph("Оглавление", style="TOC Heading")
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in heading.runs:
        set_run_font(run, 14, bold=False)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Оглавление обновится при открытии документа в Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    doc.add_page_break()


def add_paragraph(doc, text="", *, bold_lead=None, align=None, keep=False):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 14, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, 14)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 14)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.keep_together = keep
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    for run in paragraph.runs:
        set_run_font(run, 14, bold=True)
    return paragraph


def add_equation(doc, text):
    paragraph = doc.add_paragraph(style="Equation")
    run = paragraph.add_run(text)
    set_run_font(run, 12, italic=True)
    return paragraph


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = 0
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "80")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_caption(doc, text, *, figure=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if figure else WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = not figure
    run = paragraph.add_run(text)
    set_run_font(run, 14)
    return paragraph


def add_table(doc, caption, headers, rows, widths, *, font_size=10.5, aligns=None):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        shade_cell(cell, "D9E2F3")
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)

    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif aligns:
                    paragraph.alignment = aligns[column_index]
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, font_size, bold=row_index == 0)
    trailing = doc.add_paragraph()
    trailing.paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, image_path, caption, number, width=6.15):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture = run.add_picture(str(image_path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", f"Рисунок {number}")
    add_caption(doc, f"Рис. {number} – {caption}", figure=True)


def fmt(value, digits=4):
    value = float(value)
    if value != 0 and (abs(value) < 0.0001 or abs(value) >= 10000):
        return f"{value:.3e}"
    return f"{value:.{digits}f}".rstrip("0").rstrip(".")


def yes_no(value):
    return "да" if str(value).lower() == "true" else "нет"


def read(name, index_col=0):
    return pd.read_csv(RESULTS / f"{name}.csv", index_col=index_col)


def add_model_section(doc, number, title, kind, equation, explanation, table_number, fig_number):
    coefficients = read(f"{kind}_coefficients")
    metrics = read(f"{kind}_metrics").iloc[0]
    groups = read(f"{kind}_residual_groups")

    add_heading(doc, f"{number} {title}", 2)
    add_paragraph(doc, explanation)
    add_paragraph(doc, "Уравнение модели:")
    add_equation(doc, equation)
    add_paragraph(
        doc,
        "Базовая категория drive-wheels — 4wd. Индикаторы I(fwd) и I(rwd) равны 1 "
        "для соответствующего типа привода и 0 в остальных случаях.",
    )

    name_map = {
        "intercept": "Свободный член",
        "curb_weight_100": "curb-weight / 100",
        "curb_weight_100_squared": "(curb-weight / 100)²",
        "drive_fwd": "I(fwd)",
        "drive_rwd": "I(rwd)",
    }
    headers = ["Параметр", "Коэфф.", "Ст. ошибка", "Статистика", "p-value", "Значим"]
    widths = [2300, 1250, 1500, 1400, 1500, 1404]
    rows = []
    for idx, row in coefficients.iterrows():
        rows.append(
            [
                name_map[idx],
                fmt(row["coefficient"], 5),
                fmt(row["standard_error"], 5),
                fmt(row["statistic"], 4),
                fmt(row["p_value"], 5),
                yes_no(row["significant_0.05"]),
            ]
        )
    add_table(
        doc,
        f"Таблица {table_number}. Коэффициенты и их значимость для модели «{title}»",
        headers,
        rows,
        widths,
        font_size=9.5,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )

    if kind == "poisson":
        metric_rows = [
            ["Число наблюдений", int(metrics["observations"])],
            ["Сходимость", yes_no(metrics["converged"])],
            ["Число итераций", int(float(metrics["iterations"]))],
            ["Логарифм правдоподобия", fmt(metrics["log_likelihood"], 4)],
            ["Девианс", fmt(metrics["deviance"], 4)],
            ["AIC", fmt(metrics["aic"], 4)],
            ["Псевдо-R² Мак-Фаддена", fmt(metrics["mcfadden_pseudo_r_squared"], 4)],
            ["p-value модели", fmt(metrics["model_p_value"], 5)],
            ["RMSE на сдвинутой шкале", fmt(metrics["rmse_shifted_scale"], 4)],
        ]
    else:
        metric_rows = [
            ["Число наблюдений", int(metrics["observations"])],
            ["R²", fmt(metrics["r_squared"], 4)],
            ["Скорректированный R²", fmt(metrics["adjusted_r_squared"], 4)],
            ["RMSE", fmt(metrics["rmse"], 4)],
            ["MAE", fmt(metrics["mae"], 4)],
            ["F-статистика", fmt(metrics["f_statistic"], 4)],
            ["p-value модели", fmt(metrics["model_p_value"], 5)],
        ]
    if kind == "poisson":
        doc.add_page_break()
    add_table(
        doc,
        f"Таблица {table_number + 1}. Показатели качества модели «{title}»",
        ["Показатель", "Значение"],
        metric_rows,
        [6000, 3354],
        font_size=11,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_figure(
        doc,
        RESULTS / f"{kind}_residuals.png",
        f"остатки и гистограмма стандартизированных остатков модели «{title}»",
        fig_number,
    )

    max_mean = groups["mean"].abs().max()
    positive_var = groups.loc[groups["var"] > 0, "var"]
    ratio = positive_var.max() / positive_var.min()
    add_paragraph(
        doc,
        f"По четырём группам предсказанных значений максимальное по модулю среднее остатка равно "
        f"{max_mean:.3f}, а отношение наибольшей групповой дисперсии к наименьшей — {ratio:.3f}. "
        "Средние заметно отклоняются от нуля, а разброс меняется между группами, поэтому условия "
        "постоянства среднего и дисперсии остатков выполняются неудовлетворительно. Полосы на графике "
        "возникают из-за дискретных уровней зависимой переменной symboling.",
    )

    return coefficients, metrics


def build_report():
    doc = Document(REFERENCE)
    replace_title_text(doc)
    clear_after_title_page(doc)
    configure_styles(doc)
    set_update_fields(doc)

    # Вторая секция сохраняет геометрию и колонтитулы образца.
    section = doc.sections[-1]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    add_toc(doc)

    add_heading(doc, "Цель работы", 1)
    add_paragraph(
        doc,
        "Изучить построение линейной, пуассоновской и нелинейной регрессионных моделей "
        "с фиктивными переменными; провести анализ остатков; оценить качество и статистическую "
        "значимость моделей и их коэффициентов.",
    )

    add_heading(doc, "1 Описание исходных данных", 1)
    add_paragraph(
        doc,
        "В работе использован набор данных 1985 Auto Imports («Automobile»), содержащий сведения "
        "о 205 автомобилях. Для анализа выбраны три признака:",
    )
    add_paragraph(
        doc,
        "symboling — зависимая порядковая категориальная переменная, оценка страхового риска от -3 до 3; "
        "фактически наблюдаются уровни от -2 до 3.",
        bold_lead="symboling",
    )
    add_paragraph(
        doc,
        "curb-weight — независимая количественная переменная, снаряжённая масса автомобиля в фунтах.",
        bold_lead="curb-weight",
    )
    add_paragraph(
        doc,
        "drive-wheels — независимая номинальная переменная, тип привода: 4wd, fwd или rwd.",
        bold_lead="drive-wheels",
    )

    quality = read("data_quality")
    quality_rows = []
    type_map = {"ordinal": "порядковый", "nominal": "номинальный", "quantitative": "количественный"}
    for idx, row in quality.iterrows():
        quality_rows.append([idx, type_map[row["type"]], int(row["count"]), int(row["missing"]), int(row["unique"])])
    add_table(
        doc,
        "Таблица 1. Проверка выбранных данных",
        ["Признак", "Тип", "Всего", "Пропуски", "Уникальных"],
        quality_rows,
        [1900, 2400, 1500, 1700, 1854],
        font_size=10.5,
    )
    add_paragraph(doc, "Пропуски в выбранных переменных отсутствуют, поэтому все 205 наблюдений включены в моделирование.")

    doc.add_page_break()
    add_heading(doc, "2 Результаты дескриптивного анализа", 1)
    descriptive = read("descriptive_statistics").loc["curb-weight"]
    desc_rows = [
        ["Число наблюдений", int(descriptive["count"])],
        ["Среднее", fmt(descriptive["mean"], 2)],
        ["Медиана", fmt(descriptive["median"], 0)],
        ["Мода", fmt(descriptive["mode"], 0)],
        ["Минимум", fmt(descriptive["minimum"], 0)],
        ["Максимум", fmt(descriptive["maximum"], 0)],
        ["Размах", fmt(descriptive["range"], 0)],
        ["Дисперсия", fmt(descriptive["variance"], 2)],
        ["Стандартное отклонение", fmt(descriptive["standard_deviation"], 2)],
    ]
    add_table(
        doc,
        "Таблица 2. Дескриптивные показатели curb-weight",
        ["Показатель", "Значение"],
        desc_rows,
        [6000, 3354],
        font_size=11,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_paragraph(
        doc,
        "Дисперсия 271 107,87 выглядит большой из-за квадратной единицы измерения (фунт²). "
        "Для интерпретации разброса удобнее стандартное отклонение 520,68 фунта: типичная масса "
        "отклоняется от среднего примерно на 521 фунт.",
    )

    symboling = read("frequencies_symboling")
    symboling_rows = [[idx, int(row["count"]), f"{row['percent']:.2f}"] for idx, row in symboling.iterrows()]
    add_table(
        doc,
        "Таблица 3. Частоты категорий symboling",
        ["symboling", "Количество", "Доля, %"],
        symboling_rows,
        [2400, 3000, 3954],
        font_size=11,
    )
    add_paragraph(doc, "Наиболее частая категория symboling — 0: 67 автомобилей, или 32,68% выборки.")

    drive = read("frequencies_drive-wheels")
    drive_rows = [[idx, int(row["count"]), f"{row['percent']:.2f}"] for idx, row in drive.iterrows()]
    doc.add_page_break()
    add_table(
        doc,
        "Таблица 4. Частоты категорий drive-wheels",
        ["drive-wheels", "Количество", "Доля, %"],
        drive_rows,
        [3000, 3000, 3354],
        font_size=11,
    )
    add_paragraph(doc, "В выборке преобладают автомобили с передним приводом fwd: 120 наблюдений, или 58,54%.")

    add_figure(
        doc,
        RESULTS / "selected_variables_histograms.png",
        "распределения выбранных переменных",
        1,
    )
    add_paragraph(
        doc,
        "Для curb-weight число интервалов рассчитано по формуле Стерджесса: "
        "k = ceil(1 + log₂(205)) = 9. Длина интервала h = (4066 - 1488) / 9 = 286,44 фунта. "
        "Согласованность с нормальным распределением оценивается визуально. Гистограмма массы "
        "несимметрична и имеет вытянутый правый хвост, поэтому распределение curb-weight нельзя "
        "считать хорошо согласующимся с нормальным. Для категориальных symboling и drive-wheels "
        "проверка нормальности неприменима.",
    )

    add_heading(doc, "3 Регрессионные модели", 1)
    add_paragraph(
        doc,
        "Во всех моделях масса разделена на 100, чтобы коэффициент показывал изменение результата "
        "при увеличении массы на 100 фунтов. Категориальный признак drive-wheels преобразован в две "
        "фиктивные переменные; базовой категорией принят полный привод 4wd. Уровень значимости α = 0,05.",
    )

    doc.add_page_break()
    linear_coeff, linear_metrics = add_model_section(
        doc,
        "3.1",
        "Линейная регрессия",
        "linear",
        "ŷ = 2,383559 − 0,074321·(curb-weight/100) + 0,241036·I(fwd) + 0,563244·I(rwd)",
        "Модель метода наименьших квадратов рассматривает порядковые уровни symboling как числовую шкалу. "
        "Это учебное приближение, позволяющее выполнить требуемую линейную регрессию с фиктивными переменными.",
        5,
        2,
    )
    add_paragraph(
        doc,
        "Уравнение модели статистически значимо: F = 4,5703, p = 0,00404 < 0,05. "
        "Значим коэффициент curb-weight (p = 0,00105): при увеличении массы на 100 фунтов ожидаемое "
        "значение symboling уменьшается примерно на 0,074. Различия fwd и rwd относительно 4wd "
        "статистически не подтверждены. R² = 0,0639, то есть модель объясняет около 6,4% разброса результата.",
    )

    poisson_coeff, poisson_metrics = add_model_section(
        doc,
        "3.2",
        "Пуассоновская регрессия",
        "poisson",
        "E(symboling + 2 | X) = exp(1,599942 − 0,027193·(curb-weight/100) + 0,090828·I(fwd) + 0,207961·I(rwd))",
        "Распределение Пуассона требует неотрицательного отклика, поэтому использован сдвиг symboling + 2, "
        "дающий значения от 0 до 5. Поскольку symboling является порядковой, а не счётной переменной, "
        "пуассоновская модель имеет учебный характер и требует осторожной интерпретации.",
        7,
        3,
    )
    add_paragraph(
        doc,
        "Модель сошлась за 5 итераций, однако в целом незначима при α = 0,05: p = 0,06262. "
        "Коэффициент массы значим (p = 0,01603); exp(−0,027193) = 0,9732 означает уменьшение "
        "ожидаемого значения сдвинутого отклика примерно на 2,68% при увеличении массы на 100 фунтов. "
        "Коэффициенты типов привода незначимы. Псевдо-R² Мак-Фаддена равен 0,0104 и указывает на "
        "очень слабую объясняющую способность.",
    )

    nonlinear_coeff, nonlinear_metrics = add_model_section(
        doc,
        "3.3",
        "Нелинейная регрессия",
        "nonlinear",
        "ŷ = −0,269852 + 0,120912·(curb-weight/100) − 0,003537·(curb-weight/100)² + 0,325904·I(fwd) + 0,616819·I(rwd)",
        "К линейной модели добавлен квадрат нормированной массы. Модель остаётся линейной по оцениваемым "
        "коэффициентам, но описывает криволинейную зависимость результата от curb-weight.",
        9,
        4,
    )
    add_paragraph(
        doc,
        "Уравнение в целом значимо: F = 3,8981, p = 0,00452. Однако квадрат массы незначим "
        "(p = 0,17821), как и остальные отдельные коэффициенты. R² повышается лишь до 0,0723, "
        "а скорректированный R² — до 0,0538. Следовательно, убедительного нелинейного эффекта массы не выявлено.",
    )

    add_heading(doc, "4 Сравнительный анализ и интерпретация результатов", 1)
    comparison = read("model_comparison")
    model_names = {"linear": "Линейная", "nonlinear": "Нелинейная", "poisson": "Пуассоновская"}
    comparison_rows = []
    for idx, row in comparison.iterrows():
        adjusted = "—" if pd.isna(row["adjusted_r_squared"]) else fmt(row["adjusted_r_squared"], 4)
        aic = "—" if pd.isna(row["aic"]) else fmt(row["aic"], 3)
        comparison_rows.append(
            [
                model_names[idx],
                row["quality_measure"],
                fmt(row["quality_value"], 4),
                adjusted,
                fmt(row["rmse"], 4),
                aic,
                fmt(row["model_p_value"], 5),
            ]
        )
    add_table(
        doc,
        "Таблица 11. Сравнение регрессионных моделей",
        ["Модель", "Мера", "Значение", "Скорр. R²", "RMSE", "AIC", "p модели"],
        comparison_rows,
        [1550, 1900, 1100, 1300, 1100, 1100, 1304],
        font_size=8.5,
        aligns=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    add_paragraph(
        doc,
        "Среди моделей МНК нелинейная модель имеет минимальный RMSE (1,1965) и немного больший "
        "скорректированный R² (0,0538 против 0,0499). Прирост слишком мал, а коэффициент квадратного "
        "члена незначим, поэтому усложнение модели не даёт содержательного улучшения. Для практической "
        "интерпретации предпочтительнее более простая линейная модель.",
    )
    add_paragraph(
        doc,
        "Во всех вариантах объясняющая способность низкая: масса и тип привода описывают лишь малую "
        "часть различий symboling. Графики остатков содержат структуру, а групповые дисперсии отличаются "
        "в несколько раз. Это означает, что выбранные признаки и формы моделей не полностью описывают "
        "механизм страховой оценки.",
    )
    add_paragraph(
        doc,
        "Итоговый вывод: увеличение curb-weight связано со снижением оценки страхового риска, но эффект "
        "невелик. После учёта массы различия между типами привода статистически не подтверждаются. "
        "Полученные зависимости являются статистическими ассоциациями и сами по себе не доказывают "
        "причинного влияния массы или привода. Поскольку symboling — порядковая переменная, для строгого "
        "прикладного моделирования естественнее использовать порядковую логистическую регрессию; в данной "
        "работе линейная, пуассоновская и квадратичная модели построены в соответствии с заданием.",
    )

    # Чистые свойства документа без изменения пользовательских данных титульного листа.
    doc.core_properties.title = "Лабораторная работа 1.2. Регрессионный анализ"
    doc.core_properties.subject = "Методы анализа данных"
    doc.core_properties.keywords = "регрессия, автомобили, symboling, curb-weight, drive-wheels"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_report()
    print(output)

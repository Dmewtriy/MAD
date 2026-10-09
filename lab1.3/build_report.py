"""Собрать согласованные отчёты Markdown и Word из результатов ЛР 1.3.

Титульный лист сохраняется из актуального отчёта 1.2, включая исполнителей.
Статистические расчёты выполняет anova_analysis.py.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
REFERENCE = BASE.parent / "lab1.2/Отчёт_ЛР1.2_Регрессионный_анализ.docx"
OUTPUT_NAME = "Отчёт_ЛР1.3_Дисперсионный_анализ.docx"
FONT = "Times New Roman"
SOURCES = [
    ("Задание «Лабораторная работа 1.3. Дисперсионный анализ»", "ЛР1.3.pdf"),
    ("Описание исходного набора 1985 Auto Imports", "../lab1/data/imports-85.names"),
    ("SciPy: однофакторный анализ Фишера и Уэлча", "https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.f_oneway.html"),
    ("SciPy: критерий Шапиро–Уилка", "https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.shapiro.html"),
    ("SciPy: критерий Левена", "https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.levene.html"),
    ("SciPy: критерий Краскела–Уоллиса", "https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.kruskal.html"),
    ("statsmodels: таблицы ANOVA и робастная ковариация HC3", "https://www.statsmodels.org/stable/generated/statsmodels.stats.anova.anova_lm.html"),
    ("statsmodels: поправки на множественные сравнения", "https://www.statsmodels.org/stable/generated/statsmodels.stats.multitest.multipletests.html"),
]


def fmt(value, digits=2):
    if pd.isna(value):
        return "—"
    return f"{float(value):,.{digits}f}".replace(",", " ").replace(".", ",")


def pformat(value):
    value = float(value)
    if value < 0.001:
        return f"{value:.3e}".replace(".", ",")
    return fmt(value, 5)


def blocks_for_report(results):
    summary = json.loads((results / "summary.json").read_text(encoding="utf-8"))
    one, kw, two = summary["one_way"], summary["kruskal"], summary["two_way"]
    read = lambda name: pd.read_csv(results / f"{name}.csv")
    desc = read("descriptive_statistics").set_index("group")
    normality = read("normality_shapiro")
    bins = read("histogram_bins").drop_duplicates("group")
    classic_two = read("two_way_anova_type3").set_index("effect")
    robust_two = read("two_way_anova_hc3").set_index("effect")
    cells = read("two_way_cell_statistics")
    blocks, table_number, figure_number = [], 0, 0

    def h(text, level=1):
        blocks.append({"kind": "heading", "text": text, "level": level})

    def p(text):
        blocks.append({"kind": "paragraph", "text": text})

    def eq(text):
        blocks.append({"kind": "equation", "text": text})

    def table(title, headers, rows, widths=None):
        nonlocal table_number
        table_number += 1
        blocks.append({"kind": "table", "caption": f"Таблица {table_number}. {title}",
                       "headers": headers, "rows": rows,
                       "widths": widths or [1] * len(headers)})

    def figure(filename, title):
        nonlocal figure_number
        figure_number += 1
        blocks.append({"kind": "figure", "file": filename,
                       "caption": f"Рис. {figure_number} – {title}"})

    h("Цель работы")
    p("Изучить однофакторный дисперсионный анализ: сравнить цену автомобилей с разными типами привода, "
      "проверить условия применения ANOVA и выполнить критерий Краскела–Уоллиса. "
      "В дополнительном задании провести двухфакторный анализ цены по типу привода и наддуву двигателя.")
    h("1 Описание исходных данных")
    p("Пункт 1 задания выполнен заново для выбранной зависимой переменной price. Использован тот же набор "
      "1985 Auto Imports («Automobile»), что и в лабораторных работах 1.1 и 1.2. "
      "Исходный файл lab1/data/imports-85.data содержит 205 записей и 26 признаков: технические "
      "характеристики автомобилей, цену и страховые показатели. Одна строка соответствует записи "
      "об автомобиле в каталоге; данные не являются временным рядом.")
    p("price — зависимая количественная переменная, цена автомобиля в долларах США. Для сравнения групп "
      "она используется в исходной числовой шкале. Ценовые интервалы применяются только при построении гистограмм.")
    p("drive-wheels — независимая номинальная переменная, тип привода: 4wd — полный, fwd — передний, "
      "rwd — задний. Этот фактор используется в однофакторном и двухфакторном анализе.")
    p("aspiration — второй независимый номинальный фактор для пункта 10: std — стандартное исполнение "
      "без турбонаддува, turbo — турбонаддув. Выбор обусловлен техническим смыслом признака, отсутствием "
      "пропусков и наличием всех шести сочетаний с типом привода.")
    quality = read("data_quality")
    table("Проверка выбранных переменных до и после обработки пропусков",
          ["Признак", "Тип", "Всего", "Заполнено", "Пропуски", "В анализе"],
          [[r.variable, "количественный" if r.type == "quantitative" else "номинальный",
            r.total_rows, r.valid_before, r.missing_before, r.valid_after] for r in quality.itertuples()],
          [1.6, 2, 0.8, 1.2, 1.1, 1.2])
    excluded = read("excluded_rows")
    p(f"Знак «?» в исходном файле распознаётся как пропуск. Цена отсутствует в четырёх строках "
      f"(номера строк исходного файла: {', '.join(str(int(v)) for v in excluded['source_row'])}). "
      "Эти строки исключены по правилу полных наблюдений для выбранных переменных; цена не заполняется "
      "средним значением. Во всех дальнейших расчётах используется одна и та же выборка из 201 автомобиля. "
      "Пропуски в других, неиспользуемых признаках не служат причиной исключения строк. "
      "Наблюдаемые высокие цены сохранены как часть исходных данных.")
    p("Среда выполнения — Python. Использованы pandas и NumPy для подготовки данных, SciPy для "
      "статистических критериев, statsmodels для двухфакторной модели и matplotlib для графиков. "
      "Во всех проверках уровень значимости α = 0,05.")

    h("2 Результаты дескриптивного анализа")
    p("Раздел выполняет пункты 2 и 3 задания. Все характеристики рассчитаны заново для price и "
      "сформированных групп. Дисперсия и стандартное отклонение являются выборочными (делитель n − 1).")
    h("2.1 Числовые характеристики цены", 2)
    metrics = [("Число наблюдений", "count", 0), ("Среднее, долл.", "mean", 2),
               ("Медиана, долл.", "median", 2), ("Минимум, долл.", "minimum", 0),
               ("Максимум, долл.", "maximum", 0), ("Размах, долл.", "range", 0),
               ("Дисперсия, долл.²", "variance", 2), ("Стандартное отклонение, долл.", "standard_deviation", 2),
               ("Первый квартиль, долл.", "q1", 2), ("Третий квартиль, долл.", "q3", 2),
               ("Межквартильный размах, долл.", "iqr", 2), ("Асимметрия", "skewness", 3)]
    table("Дескриптивные показатели price в целом и по типу привода",
          ["Показатель", "Вся выборка", "4wd", "fwd", "rwd"],
          [[label] + [fmt(desc.loc[group, field], digits) for group in ["all", "4wd", "fwd", "rwd"]]
           for label, field, digits in metrics], [2.4, 1.3, 1.25, 1.25, 1.25])
    p(f"Моды цены всей выборки: {desc.loc['all', 'modes']} долл. Средняя цена равна "
      f"{fmt(desc.loc['all', 'mean'])} долл., медиана — {fmt(desc.loc['all', 'median'])} долл. "
      "Среднее выше медианы, что соответствует правосторонней асимметрии. Большая величина дисперсии "
      "связана с измерением в долларах²; для оценки разброса удобнее стандартное отклонение. "
      "Цены в группе rwd имеют наибольшие среднее, медиану и стандартное отклонение.")
    frequencies = []
    for variable in ["drive-wheels", "aspiration"]:
        for r in read(f"frequencies_{variable}").itertuples():
            frequencies.append([variable, r.level, int(r.count), fmt(r.percent)])
    table("Частоты категорий факторов в анализируемой выборке",
          ["Фактор", "Категория", "Число автомобилей", "Доля, %"], frequencies, [1.7, 1, 1.5, 1.2])
    figure("category_frequencies.png", "распределения типов привода и наддува в выборке из 201 автомобиля")

    h("2.2 Гистограммы и формула Стерджесса", 2)
    eq("k = ceil(1 + log₂ n);     h = (xₘₐₓ − xₘᵢₙ) / k")
    p("Здесь k — число интервалов, округлённое вверх, h — длина интервала. Расчёт выполнен по "
      "наблюдаемому диапазону цены отдельно для всей выборки и каждой группы привода. "
      "Интервалы замкнуты слева и открыты справа; последний интервал включает максимум.")
    table("Параметры гистограмм price по формуле Стерджесса",
          ["Выборка", "n", "Минимум", "Максимум", "k", "h, долл."],
          [["вся выборка" if r.group == "all" else r.group, int(r.n), fmt(desc.loc[r.group, "minimum"], 0),
            fmt(desc.loc[r.group, "maximum"], 0), int(r.k), fmt(r.h)] for r in bins.itertuples()],
          [1.5, 0.6, 1, 1, 0.6, 1.1])
    figure("price_histograms.png", "гистограммы цены с интервалами, рассчитанными по формуле Стерджесса")
    p("По гистограммам видны асимметрия и высокие значения в правом хвосте. Номинальные факторы "
      "drive-wheels и aspiration изображены столбчатыми диаграммами частот: у их категорий нет числовой "
      "длины интервала и к ним неприменима формула Стерджесса.")

    h("2.3 Оценка нормальности и равенства дисперсий", 2)
    p("Для ANOVA существенна нормальность распределения ошибок внутри групп, поэтому проверены "
      "отдельные выборки price по приводу. Общая цена проверена дополнительно для описания данных. "
      "Использованы Q–Q-графики и критерий Шапиро–Уилка: H₀ — нормальное распределение, "
      "H₁ — отклонение от него. При p < 0,05 H₀ отвергается.")
    table("Проверка нормальности цены критерием Шапиро–Уилка",
          ["Выборка", "n", "W", "p-value", "Вывод при α = 0,05"],
          [["вся выборка" if r.group == "all" else r.group, int(r.count), fmt(r.W, 4),
            pformat(r.p_value), "нормальность отвергается" if r.reject_normality else "H₀ не отвергается"]
           for r in normality.itertuples()], [1.3, 0.5, 0.8, 1.3, 2.3])
    figure("price_qq_by_drive.png", "Q–Q-графики цены для трёх групп привода")
    p("Нормальность отвергается во всех трёх группах при α = 0,05. На Q–Q-графиках точки "
      "отклоняются от прямой, особенно в хвостах. Для 4wd вывод ограничен малым объёмом — 8 наблюдений.")
    p(f"Равенство дисперсий проверено критерием Левена с центром в медиане (вариант Брауна–Форсайта). "
      f"H₀: σ²₄wd = σ²fwd = σ²rwd. Получены F = {fmt(one['levene_F'], 4)}, "
      f"df = (2; 198), p = {pformat(one['levene_p_value'])}; H₀ отвергается. "
      f"Отношение максимальной выборочной дисперсии к минимальной равно {fmt(one['variance_ratio'], 3)}. "
      "Таким образом, обычная ANOVA нарушает условия нормальности и однородности дисперсий.")

    h("3 Правила применения методов и обоснование выбора")
    p("Пункт 4 задания. Однофакторная ANOVA Фишера применяется для сравнения средних количественного "
      "признака в нескольких независимых группах. Предполагаются независимость наблюдений, нормальность "
      "ошибок внутри групп и одинаковые генеральные дисперсии. Нормальность смеси всех групп не является "
      "отдельным обязательным условием. Неравные размеры групп допустимы, но в сочетании с разными "
      "дисперсиями ухудшают надёжность обычного F-теста.")
    p("По условию работы выполняется классическая однофакторная ANOVA с F-статистикой Фишера. "
      "Её результат сопоставляется с ANOVA Уэлча, допускающей неодинаковые дисперсии и сравнивающей "
      "средние, и с ранговым критерием Краскела–Уоллиса. Уэлч устраняет предположение о равенстве "
      "дисперсий, но не гарантирует точность при произвольной форме распределения и малой группе 4wd.")
    p("Критерий Краскела–Уоллиса применяется к независимым группам числовых или порядковых "
      "наблюдений. Нормальность не требуется; в расчёте используются ранги и поправка на совпадающие "
      "значения. Для χ²-приближения желательно не менее пяти наблюдений в каждой группе: здесь минимум "
      "равен восьми. При одинаковой форме и разбросе распределений его можно трактовать как проверку "
      "различия положения, в частности медиан. Здесь разброс различается, поэтому основной вывод "
      "относится к распределениям и рангам, а не только к медианам.")
    p("Двухфакторная ANOVA оценивает два категориальных фактора и их взаимодействие. Для модели "
      "с взаимодействием нужны наблюдения в каждой комбинации факторов и остаточные степени свободы. "
      "При несбалансированном плане способ расчёта сумм квадратов необходимо указывать. Здесь "
      "используются частные суммы квадратов типа III и контрасты с нулевой суммой; дополнительно "
      "проводятся F-проверки с робастной ковариационной матрицей HC3 [7].")
    p("Независимость рассматривается как рабочее предположение: запись входит в выборку один раз "
      "и относится только к одной группе. По таблице нельзя доказать независимость; автомобили "
      "одной марки могут иметь сходные характеристики. Выборка каталога не является подтверждённой "
      "случайной выборкой рынка, поэтому выводы прежде всего описывают исследуемые данные.")

    h("4 Формулировка гипотез")
    p("Пункт 5 задания. Для однофакторного анализа цены формулируются гипотезы:")
    eq("H₀: μ₄wd = μfwd = μrwd;     H₁: хотя бы одно среднее отличается.")
    p("μ — генеральное математическое ожидание цены для соответствующего типа привода. "
      "H₁ не утверждает, что различаются все пары групп. Правило решения: отвергнуть H₀, если "
      "p < α = 0,05, или, эквивалентно, если Fнабл > Fкрит. Если p ≥ α, оснований отвергать H₀ "
      "недостаточно; это не доказывает равенство средних.")
    p("Для Краскела–Уоллиса H₀: распределения цены в группах одинаковы; H₁: хотя бы одна группа "
      "имеет отличающееся распределение. При совпадающих формах распределений это становится "
      "гипотезой об одинаковом положении. Правило решения: p < 0,05 или Hнабл > χ²крит(2).")
    p("Для двухфакторной модели отдельно проверяются H₀A — отсутствие эффекта типа привода, "
      "H₀B — отсутствие эффекта наддува и H₀AB — отсутствие взаимодействия. Альтернативы "
      "предполагают наличие соответствующего эффекта. При контрастах с нулевой суммой эффекты "
      "A и B относятся к средним, усреднённым с одинаковыми весами по уровням другого фактора. "
      "Взаимодействие означает, что разность средних цен turbo и std зависит от типа привода.")

    h("5 Градации фактора и формирование групп")
    p("Пункт 6 задания. drive-wheels является номинальным фактором с тремя готовыми градациями. "
      "Разбиение на числовые интервалы не требуется. Для каждой категории из очищенного набора "
      "отобраны соответствующие цены: 4wd — 8 значений, fwd — 118 значений, rwd — 75 значений. "
      "Группы не пересекаются и вместе включают все 201 наблюдение. Доли рассчитаны после "
      "исключения пропусков, поэтому отличаются от частот по 205 записям в прошлых работах.")
    p("Полные числовые ряды сохранены в results/price_groups.csv: каждый столбец содержит цены "
      "одной группы. Пустые клетки в конце коротких столбцов обозначают разные размеры групп, "
      "а не пропуски цены в анализируемых наблюдениях.")
    figure("price_boxplot.png", "распределения цены по типу привода: медианы, квартильный разброс и средние")

    h("6 Результаты однофакторного дисперсионного анализа")
    h("6.1 Расчёт F-статистики Фишера", 2)
    p("Пункты 7 и 8 задания. Межгрупповая сумма квадратов отражает отклонения групповых средних "
      "от общего среднего с весами, равными размерам групп. Внутригрупповая сумма квадратов "
      "отражает разброс цен относительно среднего своей группы.")
    eq("SSмеж = Σ nⱼ(ȳⱼ − ȳ)²;     SSвнутр = ΣΣ(yᵢⱼ − ȳⱼ)²")
    eq("SSобщ = SSмеж + SSвнутр;     F = [SSмеж / (k − 1)] / [SSвнутр / (N − k)]")
    p("В данном случае N = 201, k = 3; степени свободы между группами равны 2, "
      "внутри групп — 198, общие — 200. Расчёт по суммам квадратов независимо сверяется "
      "с scipy.stats.f_oneway.")
    anova = read("one_way_anova")
    source_names = {"between": "Между группами", "within": "Внутри групп", "total": "Общая"}
    table("Однофакторная ANOVA цены по типу привода",
          ["Источник", "SS, млн долл.²", "df", "MS, млн долл.²", "F", "p-value"],
          [[source_names[r.source], fmt(r.SS / 1e6, 4), int(r.df),
            fmt(r.MS / 1e6, 4), fmt(r.F, 4), pformat(r.p_value) if pd.notna(r.p_value) else "—"]
           for r in anova.itertuples()], [1.6, 1.4, 0.4, 1.4, 0.85, 1.2])
    p(f"Fнабл = {fmt(one['F'], 4)}, Fкрит(2; 198; 0,95) = {fmt(one['F_critical'], 4)}; "
      f"p = {pformat(one['p_value'])}. Fнабл > Fкрит и p < 0,05, поэтому нулевая гипотеза "
      "равенства всех средних отвергается. С учётом нарушенных предпосылок этот вывод "
      "необходимо проверить устойчивыми методами.")
    p(f"Доля межгрупповой вариации η² = SSмеж / SSобщ = {fmt(one['eta_squared'], 4)} "
      f"({fmt(one['eta_squared'] * 100)}% общего разброса цены). Скорректированная оценка "
      f"ω² = [SSмеж − (k − 1)MSвнутр] / [SSобщ + MSвнутр] = {fmt(one['omega_squared'], 4)}. "
      "Это характеристики связи в выборке, а не доля доказанного причинного влияния привода.")

    h("6.2 Проверка устойчивости: ANOVA Уэлча", 2)
    welch = read("welch_anova").iloc[0]
    table("Однофакторный анализ Уэлча при неодинаковых дисперсиях",
          ["F Уэлча", "df₁", "df₂", "Fкрит", "p-value"],
          [[fmt(welch.F, 4), int(welch.df1), fmt(welch.df2, 4), fmt(welch.F_critical, 4), pformat(welch.p_value)]])
    p(f"Уэлч также отвергает H₀: p = {pformat(one['welch_p_value'])} < 0,05. "
      "Следовательно, вывод о различии средних сохраняется при отказе от предположения "
      "об одинаковых дисперсиях.")
    pairs = read("pairwise_welch_holm")
    table("Попарные двусторонние t-тесты Уэлча с поправкой Холма",
          ["Пара", "Разность средних, долл.", "t", "df", "p Холма", "Значимо"],
          [[f"{r.group1} − {r.group2}", fmt(r.mean_difference), fmt(r.t, 4), fmt(r.df, 3),
            pformat(r.p_holm), "да" if r.reject_holm else "нет"] for r in pairs.itertuples()],
          [1.15, 1.65, 0.9, 0.9, 1.35, 0.8])
    p("Попарные сравнения выполнены дополнительно, чтобы определить источник различий. Поправка "
      "Холма контролирует ошибку первого рода в семействе трёх сравнений. Средняя цена rwd выше "
      f"средней цены 4wd на {fmt(desc.loc['rwd', 'mean'] - desc.loc['4wd', 'mean'])} долл. "
      f"и fwd на {fmt(desc.loc['rwd', 'mean'] - desc.loc['fwd', 'mean'])} долл.; оба различия "
      "значимы после поправки. Различие 4wd и fwd незначимо. Это не означает, что их "
      "генеральные средние доказанно равны.")

    h("7 Результаты критерия Краскела–Уоллиса")
    p("Пункт 9 задания. Все 201 значение цены ранжировано совместно от меньшего к большему. "
      "Совпадающим ценам назначены средние ранги; затем рассчитаны суммы рангов по группам.")
    eq("H* = [12 / (N(N + 1))] Σ(Rⱼ² / nⱼ) − 3(N + 1)")
    eq("C = 1 − Σ(t³ − t)/(N³ − N);     H = H* / C")
    p("Rⱼ — сумма рангов группы j, t — число повторений одной цены, C — поправка на "
      "совпадающие значения. Статистика H сравнивается с распределением χ² с k − 1 = 2 "
      "степенями свободы. Ручной расчёт с поправкой сверяется с scipy.stats.kruskal.")
    ranks = read("kruskal_ranks")
    table("Ранги цены по типу привода",
          ["Тип привода", "n", "Сумма рангов", "Средний ранг"],
          [[r.group, int(r.count), fmt(r.rank_sum, 1), fmt(r.mean_rank, 3)] for r in ranks.itertuples()])
    table("Критерий Краскела–Уоллиса с поправкой на совпадения",
          ["H*", "C", "H", "df", "χ²крит", "p-value"],
          [[fmt(kw['H_before_tie_correction'], 4), fmt(kw['tie_correction'], 6), fmt(kw['H'], 4),
            kw['df'], fmt(kw['chi2_critical'], 4), pformat(kw['p_value'])]], [1, 1.1, 1, 0.4, 0.9, 1.5])
    p(f"H = {fmt(kw['H'], 4)} > χ²крит = {fmt(kw['chi2_critical'], 4)}; "
      f"p = {pformat(kw['p_value'])} < 0,05. H₀ об одинаковых распределениях отвергается. "
      "Наибольший средний ранг имеет группа rwd, что соответствует более высоким ценам "
      "в этой группе. Из-за неодинакового разброса результат нельзя сводить к доказательству "
      "различия только медиан или использовать как прямую проверку равенства математических ожиданий.")
    dunn = read("pairwise_dunn_holm")
    table("Попарные сравнения рангов: критерий Данна с поправкой Холма",
          ["Пара", "z", "p Холма", "Значимо при α = 0,05"],
          [[f"{r.group1} − {r.group2}", fmt(r.z, 4), pformat(r.p_holm),
            "да" if r.reject_holm else "нет"] for r in dunn.itertuples()], [1.3, 0.8, 1.5, 1.6])
    p("Ранговые попарные сравнения подтверждают отличия rwd от 4wd и fwd; отличие "
      "4wd от fwd не подтверждается. Тем самым параметрические и ранговые процедуры "
      "дают согласованную картину при различающихся проверяемых гипотезах.")

    h("8 Дополнительное задание: двухфакторный дисперсионный анализ")
    h("8.1 Факторы, план и модель", 2)
    p("Пункт 10 задания. Зависимая переменная — price. Фактор A — drive-wheels с тремя "
      "уровнями; фактор B — aspiration с двумя уровнями. Анализ выполнен на тех же 201 "
      "наблюдении. Все шесть комбинаций присутствуют, но размеры ячеек различаются.")
    table("Число наблюдений и цена для сочетаний привода и наддува",
          ["Привод", "Наддув", "n", "Средняя, долл.", "Медиана, долл.", "Ст. откл., долл."],
          [[r.drive, r.aspiration, int(r.count), fmt(r.mean), fmt(r.median), fmt(r.standard_deviation)]
           for r in cells.itertuples()], [0.8, 0.9, 0.5, 1.5, 1.5, 1.5])
    eq("priceᵢⱼₖ = μ + Aᵢ + Bⱼ + (AB)ᵢⱼ + εᵢⱼₖ")
    p("μ — общий уровень цены; Aᵢ и Bⱼ — эффекты факторов; (AB)ᵢⱼ — взаимодействие; "
      "ε — ошибка. Модель оценивается методом наименьших квадратов с формулой "
      "price ~ C(drive, Sum) * C(aspiration, Sum). Использованы контрасты с нулевой суммой, "
      "чтобы в таблице типа III главные эффекты соответствовали сравнению средних с "
      "равными весами уровней второго фактора. Частные суммы квадратов типа III в "
      "несбалансированном плане не обязаны складываться в общую сумму квадратов.")
    p(f"Матрица модели имеет полный ранг; оценены 6 параметров. Остаточные степени свободы "
      f"равны 201 − 6 = 195. R² = {fmt(two['R_squared'], 4)}, скорректированный R² = "
      f"{fmt(two['adjusted_R_squared'], 4)}. Для ячейки 4wd × turbo имеются только 2 наблюдения: "
      "это позволяет оценить полную модель, но делает выводы по этой комбинации неточными.")

    h("8.2 Классическая ANOVA типа III", 2)
    effect_names = {"drive": "Привод A", "aspiration": "Наддув B", "interaction": "Взаимодействие A × B", "residual": "Остаток"}
    rows = []
    for effect, r in classic_two.iterrows():
        rows.append([effect_names[effect], fmt(r.SS_partial / 1e6, 4), int(r.df),
                     fmt(r.MS / 1e6, 4), fmt(r.F, 4), pformat(r.p_value) if pd.notna(r.p_value) else "—"])
    table("Двухфакторная ANOVA типа III для price",
          ["Эффект", "SSчастн., млн долл.²", "df", "MS, млн долл.²", "F", "p-value"],
          rows, [1.7, 1.45, 0.4, 1.3, 0.8, 1.2])
    p(f"Эффект привода значим: F(2; 195) = {fmt(classic_two.loc['drive', 'F'], 4)}, "
      f"p = {pformat(classic_two.loc['drive', 'p_value'])}; H₀A отвергается. "
      f"Эффект наддува незначим: F(1; 195) = {fmt(classic_two.loc['aspiration', 'F'], 4)}, "
      f"p = {pformat(classic_two.loc['aspiration', 'p_value'])}; H₀B не отвергается. "
      f"Взаимодействие незначимо: F(2; 195) = {fmt(classic_two.loc['interaction', 'F'], 4)}, "
      f"p = {pformat(classic_two.loc['interaction', 'p_value'])}; H₀AB не отвергается. "
      "Отсутствие значимости не доказывает отсутствия соответствующего эффекта.")
    figure("two_way_interaction.png", "средние цены по приводу и наддуву; указаны размеры ячеек")
    p("Непараллельность линий на графике показывает различающиеся выборочные разности "
      "turbo − std, но сама по себе не доказывает взаимодействия: для этого используется "
      "его статистическая проверка.")

    h("8.3 Проверка предпосылок и робастный анализ HC3", 2)
    p(f"Для шести ячеек критерий Левена с центром в медиане даёт F(5; 195) = "
      f"{fmt(two['cell_levene_F'], 4)}, p = {pformat(two['cell_levene_p_value'])}. "
      f"Шапиро–Уилк для остатков модели: W = {fmt(two['residual_shapiro_W'], 4)}, "
      f"p = {pformat(two['residual_shapiro_p_value'])}. Условия равенства дисперсий и "
      "нормальности ошибок не выполняются. Поэтому классическая таблица дополняется "
      "проверками тех же эффектов с робастной ковариацией HC3. Для ячейки из двух наблюдений "
      "критерий Шапиро–Уилка не рассчитывается, поскольку требуется не менее трёх значений.")
    figure("two_way_diagnostics.png", "остатки и Q–Q-график двухфакторной модели")
    table("F-проверки эффектов двухфакторной модели с ковариацией HC3",
          ["Эффект", "df₁", "df₂", "F HC3", "p-value", "Значимо"],
          [[effect_names[effect], int(r.df), int(r.df_residual), fmt(r.F, 4), pformat(r.p_value),
            "да" if r.p_value < 0.05 else "нет"] for effect, r in robust_two.iterrows()],
          [1.9, 0.4, 0.4, 0.9, 1.4, 0.7])
    p(f"Вывод о приводе сохраняется при HC3: p = {pformat(robust_two.loc['drive', 'p_value'])}. "
      f"Взаимодействие также не подтверждается: p = {pformat(robust_two.loc['interaction', 'p_value'])}. "
      f"Для наддува HC3 даёт p = {pformat(robust_two.loc['aspiration', 'p_value'])} < 0,05, "
      f"тогда как классический тест даёт p = {pformat(classic_two.loc['aspiration', 'p_value'])}. "
      "Таким образом, вывод о наддуве чувствителен к учёту неодинаковых дисперсий. Его нельзя "
      "представлять как устойчиво установленный или отсутствующий эффект. HC3 использует "
      "приближённую проверку и не устраняет проблему малой ячейки 4wd × turbo и отклонений "
      "от нормальности. По главным эффектам показаны отдельные проверки при α = 0,05; "
      "в отличие от попарных сравнений, общая поправка на три эффекта здесь не применяется.")

    h("9 Интерпретация результатов и выводы")
    comparison = [
        ["ANOVA Фишера", fmt(one['F'], 4), pformat(one['p_value']), "равенство средних отвергается"],
        ["ANOVA Уэлча", fmt(one['welch_F'], 4), pformat(one['welch_p_value']), "равенство средних отвергается"],
        ["Краскел–Уоллис", fmt(kw['H'], 4), pformat(kw['p_value']), "равенство распределений отвергается"],
    ]
    table("Сопоставление однофакторных проверок",
          ["Метод", "Статистика", "p-value", "Вывод при α = 0,05"], comparison, [1.6, 0.9, 1.4, 2.2])
    p("1. Цена автомобиля price является количественной зависимой переменной. После исключения "
      "четырёх неизвестных цен выполнен новый дескриптивный анализ 201 наблюдения, построены "
      "гистограммы с интервалами Стерджесса и проверены распределения в группах.")
    p("2. В исследуемом наборе цена связана с типом привода. Обычная ANOVA выявляет различия "
      "средних; анализ Уэлча подтверждает этот вывод при неодинаковых дисперсиях. Критерий "
      "Краскела–Уоллиса независимо подтверждает различия распределений цены.")
    p(f"3. Средние цены равны: 4wd — {fmt(desc.loc['4wd', 'mean'])} долл., fwd — "
      f"{fmt(desc.loc['fwd', 'mean'])} долл., rwd — {fmt(desc.loc['rwd', 'mean'])} долл. "
      "Попарные проверки Уэлча и Данна с поправкой Холма показывают, что rwd отличается "
      "от двух других групп. Различие 4wd и fwd статистически не подтверждено.")
    p("4. В двухфакторной модели с наддувом значимость привода сохраняется, а взаимодействие "
      "не подтверждено ни обычной, ни робастной проверкой. Вывод о наддуве зависит от метода: "
      "классическая ANOVA типа III не отвергает H₀B, HC3 отвергает её при 0,05. Результат "
      "по этому фактору требует осторожной интерпретации из-за несбалансированного плана, "
      "неодинаковых дисперсий и малой ячейки.")
    p("5. Выявленные различия не доказывают причинного влияния привода или наддува на цену. "
      "Группы могут различаться по марке, классу, мощности и комплектации; эти признаки не "
      "контролировались. Результаты относятся к историческим данным каталога, а не к "
      "актуальным рыночным ценам автомобилей.")

    h("Список использованных материалов")
    for i, (label, link) in enumerate(SOURCES, 1):
        blocks.append({"kind": "source", "text": f"{i}. {label}", "url": link})
    h("Приложение. Воспроизведение расчётов")
    p("Расчёты выполняются файлом anova_analysis.py. Отчёты REPORT.md и "
      "Отчёт_ЛР1.3_Дисперсионный_анализ.docx собираются файлом build_report.py "
      "из сохранённых CSV, PNG и summary.json. Команды запуска приведены в README.md. "
      "Файл results/analysis_data.csv содержит очищенные данные с номерами исходных строк; "
      "results/excluded_rows.csv — исключённые строки.")
    p("Контроль вычислений: суммы квадратов однофакторной модели сверены с общим разбросом; "
      "F Фишера, F Уэлча и H Краскела–Уоллиса независимо рассчитаны по формулам и сопоставлены "
      "с SciPy. Частные F двухфакторной модели сверены с вложенными моделями, оценёнными "
      "через numpy.linalg.lstsq. Все проверки выполнены успешно.")
    return blocks


def set_font(run, size=14, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    for name in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn(f"w:{name}"), FONT)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(14)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = pf.space_after = Pt(0)
    for name in ["Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.bold = FONT, Pt(14), True
        style.font.color.rgb = RGBColor(0, 0, 0)
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = pf.left_indent = Cm(0)
        pf.space_before, pf.space_after = Pt(10), Pt(4)
        pf.line_spacing, pf.keep_with_next = 1.0, True


def prepare_title(reference):
    doc = Document(reference)
    body = doc._element.body
    split = next((child for child in body if child.tag == qn("w:p")
                  and child.find(".//" + qn("w:sectPr")) is not None), None)
    if split is None:
        raise RuntimeError("В отчёте-образце не найден разрыв после титульного листа")
    after = False
    for child in list(body):
        if after and child.tag != qn("w:sectPr"):
            body.remove(child)
        if child is split:
            after = True
    # После удаления основного текста убираем связи с рисунками прошлой работы.
    # Иначе Word сохраняет их внутри архива, хотя в новом отчёте они не видны.
    used_images = set(doc._element.xpath("//@r:embed"))
    for relation_id, relation in list(doc.part.rels.items()):
        if relation.reltype.endswith("/image") and relation_id not in used_images:
            doc.part.drop_rel(relation_id)
    for paragraph in doc.paragraphs:
        if "ЛАБОРАТОРНАЯ РАБОТА №" in paragraph.text:
            paragraph.clear()
            set_font(paragraph.add_run("ЛАБОРАТОРНАЯ РАБОТА №1.3"), bold=True)
        elif "Регрессионный анализ" in paragraph.text:
            paragraph.clear()
            first = paragraph.add_run("«Дисперсионный анализ»")
            set_font(first, bold=True)
            first.add_break()
            set_font(paragraph.add_run("по дисциплине «Методы анализа данных»"), size=12)
    configure_styles(doc)
    section = doc.sections[-1]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)
    doc.core_properties.title = "Лабораторная работа 1.3. Дисперсионный анализ"
    doc.core_properties.subject = "Цена автомобилей Automobile по приводу и наддуву"
    return doc


def paragraph(doc, text, *, caption=False, centered=False, keep_next=False):
    p = doc.add_paragraph()
    set_font(p.add_run(text))
    if caption or centered:
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.keep_with_next = keep_next
    return p


def add_word_table(doc, block):
    paragraph(doc, block["caption"], caption=True, keep_next=True)
    table = doc.add_table(rows=1, cols=len(block["headers"]))
    table.style = "Table Grid"
    table.autofit = False
    # 16,5 см — доступная ширина страницы A4 с полями 3 и 1,5 см.
    fractions = block["widths"]
    widths = [Cm(16.5 * weight / sum(fractions)) for weight in fractions]
    for column, width in zip(table.columns, widths):
        column.width = width
    for cell, header in zip(table.rows[0].cells, block["headers"]):
        cell.text = str(header)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9E2F3")
        cell._tc.get_or_add_tcPr().append(shading)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in block["rows"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.space_before = pf.space_after = Pt(2)
                pf.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 and row_index else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, size=10.5 if len(widths) >= 5 else 11, bold=row_index == 0)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.5
    return table


def write_word(blocks, results, output, reference):
    doc = prepare_title(reference)
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            p = doc.add_paragraph(block["text"], style=f"Heading {block['level']}")
            for run in p.runs:
                set_font(run, bold=True)
        elif kind == "paragraph":
            paragraph(doc, block["text"])
        elif kind == "equation":
            p = paragraph(doc, block["text"], centered=True)
            p.paragraph_format.keep_together = True
            for run in p.runs:
                set_font(run, size=12, italic=True)
        elif kind == "table":
            add_word_table(doc, block)
        elif kind == "figure":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            picture = p.add_run().add_picture(str(results / block["file"]), width=Cm(16.3))
            picture._inline.docPr.set("descr", block["caption"])
            paragraph(doc, block["caption"], caption=True, centered=True)
        elif kind == "source":
            p = paragraph(doc, f"{block['text']}. {block['url']}")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_font(run, size=12)
    path = output / OUTPUT_NAME
    doc.save(path)
    return path


def write_markdown(blocks, results, output):
    lines = ["# Лабораторная работа 1.3. Дисперсионный анализ", ""]
    for b in blocks:
        if b["kind"] == "heading":
            lines += ["#" * (b["level"] + 1) + " " + b["text"], ""]
        elif b["kind"] == "paragraph":
            lines += [b["text"], ""]
        elif b["kind"] == "equation":
            lines += [f"`{b['text']}`", ""]
        elif b["kind"] == "table":
            lines += [b["caption"], "", "| " + " | ".join(b["headers"]) + " |",
                      "| " + " | ".join(["---"] * len(b["headers"])) + " |"]
            lines += ["| " + " | ".join(str(v).replace("|", "\\|") for v in row) + " |" for row in b["rows"]]
            lines.append("")
        elif b["kind"] == "figure":
            # При обычном запуске results находится рядом с Markdown-отчётом.
            import os
            relative = Path(os.path.relpath(results / b["file"], output)).as_posix()
            lines += [f"![{b['caption']}]({relative})", ""]
        elif b["kind"] == "source":
            lines += [f"{b['text']}. [{b['url']}]({b['url']})", ""]
    path = output / "REPORT.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--results", type=Path, default=BASE / "results")
    parser.add_argument("--output", type=Path, default=BASE)
    parser.add_argument("--reference", type=Path, default=REFERENCE)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    blocks = blocks_for_report(args.results)
    print(write_markdown(blocks, args.results, args.output))
    print(write_word(blocks, args.results, args.output, args.reference))


if __name__ == "__main__":
    main()
